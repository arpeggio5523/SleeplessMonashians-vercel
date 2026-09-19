#!/usr/bin/env python3
"""Quick rule-based baseline: classify -> extract -> normalize -> compare."""
import json, re, os, glob, subprocess, sys, shutil, unicodedata
from pathlib import Path

def _find(*names):
    for n in names:
        p = Path(n)
        if p.exists(): return p
    raise SystemExit(f"Could not find any of {names} in {Path.cwd()}")

ROOT   = _find("sdoc-hackathon-bundle", "bundle", "sdoc_hackathon_bundle")
DOCKER = _find("sdoc-hackathon-docker", "docker", "sdoc_hackathon_docker")
FIELDS = ["shipper","consignee","notify_party","port_of_loading",
          "port_of_discharge","container_count","gross_weight_kg"]

LABELS = {
    "shipper": ["Shipper/Exporter","Shipper (Principal or Seller)","Shipper","SHIPPER"],
    "consignee": ["Consignee (Non-Negotiable)","To the Order of","Consignee","CONSIGNEE"],
    "notify_party": ["Notify Party/Intermediate Consignee","Notify Party","Notify","NOTIFY PARTY"],
    "port_of_loading": ["Port of Loading (POL)","PORT OF LOADING","Port of Loading","Load Port","POL"],
    "port_of_discharge": ["Port of Discharge (POD)","PORT OF DISCHARGE","Port of Discharge","Discharge Port","POD"],
    "container_count": ["No. of Containers or Packages","No. of Containers","Total Containers","Container Count"],
    "gross_weight_kg": ["TOTAL Gross Weight (KG)","TOTAL Gross Wt (kgs)","Gross Weight (KG)","Gross Wt (kgs)","Gross Weight毛重(KGS)","GROSS WEIGHT","Gross Weight"],
}

# ---------- text extraction ----------
def read_text(path):
    p = ROOT / path
    if not p.exists() or p.stat().st_size == 0:
        return None
    ext = p.suffix.lower()
    try:
        if ext == ".txt":
            raw = p.read_bytes()
            for enc in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
                try:
                    return raw.decode(enc)
                except UnicodeDecodeError:
                    continue
            return raw.decode("utf-8", "replace")
        if ext == ".pdf":
            t = None
            if shutil.which("pdftotext"):                 # fastest when present
                out = subprocess.run(["pdftotext","-layout",str(p),"-"],
                                     capture_output=True, timeout=30)
                t = out.stdout.decode("utf-8","replace")
            if not t or len(t.strip()) <= 40:             # Windows / no poppler
                try:
                    import pdfplumber
                    with pdfplumber.open(str(p)) as pdf:
                        t = "\n".join((pg.extract_text(layout=True) or "")
                                       for pg in pdf.pages)
                except Exception:
                    t = None
            return t if t and len(t.strip()) > 40 else None
        if ext == ".docx":
            import docx
            d = docx.Document(str(p)); parts=[]
            for para in d.paragraphs:
                if para.text.strip(): parts.append(para.text)
            for tb in d.tables:
                for row in tb.rows:
                    cells=[c.text.replace("\n"," | ") for c in row.cells]
                    parts.append(": ".join(cells))
            return "\n".join(parts)
        if ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(str(p), data_only=True); parts=[]
            for ws in wb:
                for r in ws.iter_rows(values_only=True):
                    vals=[str(c) for c in r if c is not None]
                    if vals: parts.append(": ".join(vals))
            return "\n".join(parts)
    except Exception:
        return None
    return None

DOC_MARKERS = {
    "SI": ["SHIPPING INSTRUCTION","BL INSTRUCTION","BILL OF LADING INSTRUCTION"],
    "BL": ["BILL OF LADING"],
    "WRONG": ["COMMERCIAL INVOICE","PACKING LIST","CERTIFICATE OF ORIGIN"],
}
def doc_kind(text):
    head = text[:400].upper()
    for m in DOC_MARKERS["WRONG"]:
        if m in head: return "WRONG"
    return "DOC"

# ---------- field extraction ----------
BLANKS = {"???","_______","TBA","TBC","","N/A","____MT","____","N.A."}
def clean_label_noise(s):
    # Drop any non-ASCII (CJK labels, or cp1252 mojibake of them) so that
    # "Gross Weight<CJK>(KGS):" and "Gross Weight<junk>(KGS):" both match.
    s = unicodedata.normalize("NFKC", s)
    s = "".join(ch if ord(ch) < 128 else "" for ch in s)
    return s.strip()

def extract(text):
    out = {}
    lines = [l.rstrip() for l in text.split("\n")]
    for field, labs in LABELS.items():
        for lab in labs:
            pat = re.compile(r"^\s*(?:TOTAL\s+)?" + re.escape(lab) + r"\s*(?:\([^)]*\))?\s*[::]\s*(.+)$", re.I)
            hit = None
            for l in lines:
                l2 = clean_label_noise(l)
                m = pat.match(l2)
                if m: hit = m.group(1).strip(); break
            if hit is None:  # whitespace-column layout (pdf/docx tables)
                pat2 = re.compile(r"^\s*(?:TOTAL\s+)?" + re.escape(lab) + r"\s*(?:\([^)]*\))?\s{2,}(.+)$", re.I)
                for l in lines:
                    l2 = clean_label_noise(l)
                    m = pat2.match(l2)
                    if m: hit = m.group(1).strip(); break
            if hit is None:  # last resort: single-space separated (pdfplumber)
                pat3 = re.compile(r"^\s*(?:TOTAL\s+)?" + re.escape(lab) +
                                  r"\s*(?:\([^)]*\))?\s+(\S.*)$", re.I)
                for l in lines:
                    l2 = clean_label_noise(l)
                    m = pat3.match(l2)
                    if m and not m.group(1).strip().startswith(("|",":")):
                        hit = m.group(1).strip(); break
            if hit is not None:
                out[field] = hit
                break
    return out

# ---------- normalization ----------
def norm_party(v):
    if v is None: return None
    v = v.split("|")[0]
    v = v.split("\n")[0]
    v = v.upper()
    v = re.sub(r"[.,]", " ", v)
    v = re.sub(r"\b(CO|COMPANY|LIMITED|LTD|PVT|PTE|SDN|BHD|LLC|INC|GMBH|FZE|FZ|CORP|PTY)\b", " ", v)
    v = re.sub(r"[^A-Z0-9 ]", " ", v)
    return " ".join(v.split()) or None

def norm_port(v):
    if v is None: return None
    v = re.sub(r"\([A-Z]{5}\)", "", v)        # drop UN/LOCODE -- name is authoritative
    v = v.upper()
    v = re.sub(r"[^A-Z0-9 ]", " ", v)
    return " ".join(v.split()) or None

def norm_count(v):
    if v is None: return None
    m = re.search(r"(\d+)", str(v).replace(",",""))
    return int(m.group(1)) if m else None

def norm_weight(v):
    if v is None: return None
    s = str(v).upper().replace(",","")
    m = re.search(r"(\d+(?:\.\d+)?)", s)
    if not m: return None
    n = float(m.group(1))
    if "LB" in s: n *= 0.45359237
    if re.search(r"\bMT\b|TONNE|TON\b", s): n *= 1000
    return round(n)

NORM = {"shipper":norm_party,"consignee":norm_party,"notify_party":norm_party,
        "port_of_loading":norm_port,"port_of_discharge":norm_port,
        "container_count":norm_count,"gross_weight_kg":norm_weight}

def is_blank(raw):
    if raw is None: return True
    return raw.strip().strip("_").strip() in BLANKS or raw.strip() in BLANKS

# ---------- classification ----------
def classify(e):
    b = " ".join((e.get("body") or "").split()).upper()
    s = (e.get("subject") or "").upper()
    # BODY first -- subjects are deliberately misleading
    if "PLEASE COMPARE THE SI AND DRAFT BL" in b or "SEND THE DRAFT BL" in b \
       or "DRAFT BL FOR" in b:
        return "BL_COMPARISON"
    if "PLEASE FIND SHIPPING INSTRUCTION" in b or "SHIPPING INSTRUCTION FOR" in b:
        return "SI_REQUEST"
    if "QUERY ON INVOICE" in b or "D&D" in b or "DETENTION CHARGES" in b \
       or "GR IS STILL MISSING" in b or "INVOICE" in b:
        return "INVOICE_QUERY"
    if any(k in b for k in ["OUTSTANDING BL","BERTHING REPORT","SUBMIT SI & AED",
                            "UPDATE SUMMARY","LOADING COMPLETED","PENDING ITEMS"]):
        return "GENERAL"
    if any(k in b for k in ["CONGRATULAT","PRIZE","CLAIM YOUR","MAILBOX HAS EXCEEDED",
                            "VERIFY YOUR ACCOUNT","LIMITED TIME OFFER","UNPAID CUSTOMS",
                            "BITCOIN","GIFT CARD","CLICK HERE","HTTP://","GUARANTEED"]):
        return "SPAM"
    if e.get("attachments"): return "BL_COMPARISON"
    return "GENERAL"

# ---------- main ----------
def main():
    emails = [json.loads(p.read_text()) for p in sorted((ROOT/"inbox").glob("email_*.json"))]
    sub = {}
    for e in emails:
        eid = e["email_id"]
        cat = classify(e)
        rec = {"category":cat,"status":"OK","review_reason":None,
               "has_defect":False,"defect_fields":[]}
        if cat == "BL_COMPARISON":
            atts = e.get("attachments") or []
            body = (e.get("body") or "").upper()
            asks_compare = "COMPARE THE SI" in body
            if len(atts) < 2:
                if asks_compare:
                    rec.update(status="NEEDS_REVIEW", review_reason="missing_attachment")
            else:
                si_t = read_text(atts[0]); bl_t = read_text(atts[1])
                if si_t is None or bl_t is None:
                    rec.update(status="NEEDS_REVIEW", review_reason="unreadable")
                elif doc_kind(si_t)=="WRONG" or doc_kind(bl_t)=="WRONG":
                    rec.update(status="NEEDS_REVIEW", review_reason="wrong_doc_type")
                else:
                    si_r, bl_r = extract(si_t), extract(bl_t)
                    missing = False; defects = []
                    for f in FIELDS:
                        a, b = si_r.get(f), bl_r.get(f)
                        if is_blank(a) or is_blank(b): missing = True; continue
                        na, nb = NORM[f](a), NORM[f](b)
                        if na is None or nb is None: missing = True; continue
                        if na != nb: defects.append(f)
                    if missing:
                        rec.update(status="NEEDS_REVIEW", review_reason="missing_value")
                    elif defects:
                        rec.update(status="MISMATCH", has_defect=True,
                                   defect_fields=sorted(defects))
        sub[eid] = rec
    json.dump(sub, open("submission.json","w"), indent=1)

    sys.path.insert(0, str(DOCKER/"server"))
    import scoring
    gt = json.load(open(DOCKER/"data_v2"/"ground_truth.json"))
    r = scoring.score_all(gt, sub)
    print(f"FINAL SCORE       {r['final_score']:.4f}")
    print(f"  stage1 macroF1  {r['stage1']['macro_f1']:.4f}   acc {r['stage1']['accuracy']:.4f}")
    print(f"  stage3 defectF1 {r['stage3']['defect_f1']:.4f}   P {r['stage3']['defect_precision']:.3f} R {r['stage3']['defect_recall']:.3f}")
    print(f"  field F1        {r['stage3']['field_f1']:.4f}   exact {r['stage3']['exact_match_rate']:.3f}")
    print(f"  end-to-end      {r['end_to_end']['rate']:.4f}  ({r['end_to_end']['success']}/{r['end_to_end']['total']})")
    print(f"  escalation      R {r['reliability']['escalation_recall']:.3f} P {r['reliability']['escalation_precision']:.3f}")
    print("\nCONFUSION (actual -> predicted):")
    for a,d in r["stage1"]["confusion"].items():
        print(f"  {a:15} {dict(d)}")
    print("\nper review reason:", r["reliability"]["per_reason"])

main()
