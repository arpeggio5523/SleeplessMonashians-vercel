"""
Core shipping-document verification logic.

This is baseline.py's classify -> extract -> normalize -> compare pipeline,
refactored to work on in-memory data (bytes / dicts) instead of reading
files off local disk, so it can be called from an API request.
"""
import re
import io
import subprocess
import shutil
import tempfile
import unicodedata
from pathlib import Path

FIELDS = ["shipper", "consignee", "notify_party", "port_of_loading",
          "port_of_discharge", "container_count", "gross_weight_kg"]

LABELS = {
    "shipper": ["Shipper/Exporter", "Shipper (Principal or Seller)", "Shipper", "SHIPPER"],
    "consignee": ["Consignee (Non-Negotiable)", "To the Order of", "Consignee", "CONSIGNEE"],
    "notify_party": ["Notify Party/Intermediate Consignee", "Notify Party", "Notify", "NOTIFY PARTY"],
    "port_of_loading": ["Port of Loading (POL)", "PORT OF LOADING", "Port of Loading", "Load Port", "POL"],
    "port_of_discharge": ["Port of Discharge (POD)", "PORT OF DISCHARGE", "Port of Discharge", "Discharge Port", "POD"],
    "container_count": ["No. of Containers or Packages", "No. of Containers", "Total Containers", "Container Count"],
    "gross_weight_kg": ["TOTAL Gross Weight (KG)", "TOTAL Gross Wt (kgs)", "Gross Weight (KG)",
                         "Gross Wt (kgs)", "Gross Weight毛重(KGS)", "GROSS WEIGHT", "Gross Weight"],
}

DOC_MARKERS_WRONG = ["COMMERCIAL INVOICE", "PACKING LIST", "CERTIFICATE OF ORIGIN"]
BLANKS = {"???", "_______", "TBA", "TBC", "", "N/A", "____MT", "____", "N.A."}


# ---------- text extraction (from bytes, not disk paths) ----------
def read_text_from_bytes(filename: str, content: bytes):
    """Extract text from an attachment given its filename (for extension) and raw bytes."""
    ext = Path(filename).suffix.lower()
    if not content:
        return None
    try:
        if ext == ".txt":
            for enc in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
                try:
                    return content.decode(enc)
                except UnicodeDecodeError:
                    continue
            return content.decode("utf-8", "replace")

        if ext == ".pdf":
            t = None
            with tempfile.NamedTemporaryFile(suffix=".pdf") as tmp:
                tmp.write(content)
                tmp.flush()
                if shutil.which("pdftotext"):
                    out = subprocess.run(["pdftotext", "-layout", tmp.name, "-"],
                                          capture_output=True, timeout=30)
                    t = out.stdout.decode("utf-8", "replace")
                if not t or len(t.strip()) <= 40:
                    try:
                        import pdfplumber
                        with pdfplumber.open(tmp.name) as pdf:
                            t = "\n".join((pg.extract_text(layout=True) or "") for pg in pdf.pages)
                    except Exception:
                        t = None
            return t if t and len(t.strip()) > 40 else None

        if ext == ".docx":
            import docx
            d = docx.Document(io.BytesIO(content))
            parts = []
            for para in d.paragraphs:
                if para.text.strip():
                    parts.append(para.text)
            for tb in d.tables:
                for row in tb.rows:
                    cells = [c.text.replace("\n", " | ") for c in row.cells]
                    parts.append(": ".join(cells))
            return "\n".join(parts)

        if ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            parts = []
            for ws in wb:
                for r in ws.iter_rows(values_only=True):
                    vals = [str(c) for c in r if c is not None]
                    if vals:
                        parts.append(": ".join(vals))
            return "\n".join(parts)
    except Exception:
        return None
    return None


def doc_kind(text):
    head = text[:400].upper()
    for m in DOC_MARKERS_WRONG:
        if m in head:
            return "WRONG"
    return "DOC"


# ---------- field extraction ----------
def clean_label_noise(s):
    s = unicodedata.normalize("NFKC", s)
    s = "".join(ch if ord(ch) < 128 else "" for ch in s)
    return s.strip()


def extract(text):
    out = {}
    lines = [l.rstrip() for l in text.split("\n")]
    for field, labs in LABELS.items():
        for lab in labs:
            pat = re.compile(r"^\s*(?:TOTAL\s+)?" + re.escape(lab) + r"\s*(?:\([^)]*\))?\s*[::]\s*(.+)$", re.I)
            hit = None
            for l in lines:
                l2 = clean_label_noise(l)
                m = pat.match(l2)
                if m:
                    hit = m.group(1).strip()
                    break
            if hit is None:
                pat2 = re.compile(r"^\s*(?:TOTAL\s+)?" + re.escape(lab) + r"\s*(?:\([^)]*\))?\s{2,}(.+)$", re.I)
                for l in lines:
                    l2 = clean_label_noise(l)
                    m = pat2.match(l2)
                    if m:
                        hit = m.group(1).strip()
                        break
            if hit is None:
                pat3 = re.compile(r"^\s*(?:TOTAL\s+)?" + re.escape(lab) +
                                   r"\s*(?:\([^)]*\))?\s+(\S.*)$", re.I)
                for l in lines:
                    l2 = clean_label_noise(l)
                    m = pat3.match(l2)
                    if m and not m.group(1).strip().startswith(("|", ":")):
                        hit = m.group(1).strip()
                        break
            if hit is not None:
                out[field] = hit
                break
    return out


# ---------- normalization ----------
def norm_party(v):
    if v is None:
        return None
    v = v.split("|")[0]
    v = v.split("\n")[0]
    v = v.upper()
    v = re.sub(r"[.,]", " ", v)
    v = re.sub(r"\b(CO|COMPANY|LIMITED|LTD|PVT|PTE|SDN|BHD|LLC|INC|GMBH|FZE|FZ|CORP|PTY)\b", " ", v)
    v = re.sub(r"[^A-Z0-9 ]", " ", v)
    return " ".join(v.split()) or None


def norm_port(v):
    if v is None:
        return None
    v = re.sub(r"\([A-Z]{5}\)", "", v)
    v = v.upper()
    v = re.sub(r"[^A-Z0-9 ]", " ", v)
    return " ".join(v.split()) or None


def norm_count(v):
    if v is None:
        return None
    m = re.search(r"(\d+)", str(v).replace(",", ""))
    return int(m.group(1)) if m else None


def norm_weight(v):
    if v is None:
        return None
    s = str(v).upper().replace(",", "")
    m = re.search(r"(\d+(?:\.\d+)?)", s)
    if not m:
        return None
    n = float(m.group(1))
    if "LB" in s:
        n *= 0.45359237
    if re.search(r"\bMT\b|TONNE|TON\b", s):
        n *= 1000
    return round(n)


NORM = {"shipper": norm_party, "consignee": norm_party, "notify_party": norm_party,
        "port_of_loading": norm_port, "port_of_discharge": norm_port,
        "container_count": norm_count, "gross_weight_kg": norm_weight}


def is_blank(raw):
    if raw is None:
        return True
    return raw.strip().strip("_").strip() in BLANKS or raw.strip() in BLANKS


# ---------- classification ----------
def classify(subject: str, body: str, has_attachments: bool):
    b = " ".join((body or "").split()).upper()
    if "PLEASE COMPARE THE SI AND DRAFT BL" in b or "SEND THE DRAFT BL" in b \
       or "DRAFT BL FOR" in b:
        return "BL_COMPARISON"
    if "PLEASE FIND SHIPPING INSTRUCTION" in b or "SHIPPING INSTRUCTION FOR" in b:
        return "SI_REQUEST"
    if "QUERY ON INVOICE" in b or "D&D" in b or "DETENTION CHARGES" in b \
       or "GR IS STILL MISSING" in b or "INVOICE" in b:
        return "INVOICE_QUERY"
    if any(k in b for k in ["OUTSTANDING BL", "BERTHING REPORT", "SUBMIT SI & AED",
                             "UPDATE SUMMARY", "LOADING COMPLETED", "PENDING ITEMS"]):
        return "GENERAL"
    if any(k in b for k in ["CONGRATULAT", "PRIZE", "CLAIM YOUR", "MAILBOX HAS EXCEEDED",
                             "VERIFY YOUR ACCOUNT", "LIMITED TIME OFFER", "UNPAID CUSTOMS",
                             "BITCOIN", "GIFT CARD", "CLICK HERE", "HTTP://", "GUARANTEED"]):
        return "SPAM"
    if has_attachments:
        return "BL_COMPARISON"
    return "GENERAL"


# ---------- full single-email pipeline ----------
def process_email(subject: str, body: str, attachments: list):
    """
    attachments: list of dicts, each {"filename": str, "content": bytes}
    Order matters for BL_COMPARISON: attachments[0] = SI, attachments[1] = BL,
    matching the convention in baseline.py.
    """
    cat = classify(subject, body, has_attachments=bool(attachments))
    rec = {
        "category": cat,
        "status": "OK",
        "review_reason": None,
        "has_defect": False,
        "defect_fields": [],
        "field_comparison": {},  # populated for BL_COMPARISON: {field: {"si": ..., "bl": ..., "match": bool}}
    }

    if cat != "BL_COMPARISON":
        return rec

    body_u = (body or "").upper()
    asks_compare = "COMPARE THE SI" in body_u

    if len(attachments) < 2:
        if asks_compare:
            rec.update(status="NEEDS_REVIEW", review_reason="missing_attachment")
        return rec

    si_t = read_text_from_bytes(attachments[0]["filename"], attachments[0]["content"])
    bl_t = read_text_from_bytes(attachments[1]["filename"], attachments[1]["content"])

    if si_t is None or bl_t is None:
        rec.update(status="NEEDS_REVIEW", review_reason="unreadable")
        return rec

    if doc_kind(si_t) == "WRONG" or doc_kind(bl_t) == "WRONG":
        rec.update(status="NEEDS_REVIEW", review_reason="wrong_doc_type")
        return rec

    si_r, bl_r = extract(si_t), extract(bl_t)
    missing = False
    defects = []
    comparison = {}
    for f in FIELDS:
        a, b = si_r.get(f), bl_r.get(f)
        if is_blank(a) or is_blank(b):
            missing = True
            continue
        na, nb = NORM[f](a), NORM[f](b)
        if na is None or nb is None:
            missing = True
            continue
        match = (na == nb)
        comparison[f] = {"si": a, "bl": b, "match": match}
        if not match:
            defects.append(f)

    rec["field_comparison"] = comparison

    if missing:
        rec.update(status="NEEDS_REVIEW", review_reason="missing_value")
    elif defects:
        rec.update(status="MISMATCH", has_defect=True, defect_fields=sorted(defects))

    return rec
